"""Simple document loaders for TXT, DOCX, and PDF (lightweight placeholders).

These are minimal implementations intended as a starting point. Replace with
LangChain loaders when wiring the full pipeline.
"""
from typing import List, Dict
from pathlib import Path
import json

try:
    import docx
except Exception:
    docx = None

try:
    from pypdf import PdfReader
except Exception:
    PdfReader = None


def find_files(directory: str, extensions=None) -> List[Path]:
    p = Path(directory)
    if extensions is None:
        extensions = [".txt", ".docx", ".pdf"]
    return [f for f in p.rglob("*") if f.suffix.lower() in extensions]


def load_txt(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")


def load_docx(path: Path) -> str:
    if docx is None:
        raise RuntimeError("python-docx not installed")
    doc = docx.Document(path)
    paragraphs = [p.text for p in doc.paragraphs]
    return "\n".join(paragraphs)


def load_pdf(path: Path) -> str:
    if PdfReader is None:
        raise RuntimeError("pypdf not installed")
    reader = PdfReader(path)
    texts = []
    for page in reader.pages:
        try:
            texts.append(page.extract_text() or "")
        except Exception:
            texts.append("")
    return "\n".join(texts)


def load_file(path: str) -> Dict:
    p = Path(path)
    if p.suffix.lower() == ".txt":
        content = load_txt(p)
    elif p.suffix.lower() == ".docx":
        content = load_docx(p)
    elif p.suffix.lower() == ".pdf":
        content = load_pdf(p)
    else:
        raise ValueError(f"Unsupported file type: {p.suffix}")
    return {"source": str(p), "text": content}


def chunk_text(text: str, chunk_size: int = 1000, overlap: int = 200) -> List[str]:
    """Very simple character-based chunking.

    Replace with sentence-based chunking (spaCy) for production.
    """
    if not text:
        return []
    chunks = []
    start = 0
    length = len(text)
    while start < length:
        end = min(start + chunk_size, length)
        chunks.append(text[start:end])
        start = end - overlap if end < length else end
    return chunks
